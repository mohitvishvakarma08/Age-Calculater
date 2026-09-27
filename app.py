from datetime import date
from io import BytesIO
import calendar

from flask import Flask, jsonify, render_template, request, send_file

app = Flask(__name__)


def calculate(birth, end):
    if end < birth:
        raise ValueError("The calculation date cannot be before the date of birth.")

    years = end.year - birth.year
    months = end.month - birth.month
    days = end.day - birth.day

    if days < 0:
        months -= 1
        previous_month = end.month - 1 or 12
        previous_year = end.year if end.month > 1 else end.year - 1
        days += calendar.monthrange(previous_year, previous_month)[1]

    if months < 0:
        years -= 1
        months += 12

    elapsed_days = (end - birth).days
    next_year = end.year if (end.month, end.day) < (birth.month, birth.day) else end.year + 1
    birthday_day = min(birth.day, calendar.monthrange(next_year, birth.month)[1])
    next_birthday = date(next_year, birth.month, birthday_day)
    days_until = (next_birthday - end).days

    return {
        "age": {"years": years, "months": months, "days": days},
        "born_on": birth.strftime("%A"),
        "next_birthday": f"{days_until} days",
        "birthday_date": next_birthday.strftime("%B %d, %Y"),
        "totals": {
            "months": years * 12 + months,
            "weeks": elapsed_days // 7,
            "days": elapsed_days,
            "hours": elapsed_days * 24,
            "minutes": elapsed_days * 24 * 60,
        },
    }


@app.route("/")
def index():
    return render_template("index.html", today=date.today().isoformat())


@app.post("/api/calculate")
def calculate_route():
    try:
        data = request.get_json()
        birth = date.fromisoformat(data["birth"])
        end = date.fromisoformat(data["end"])
        result = calculate(birth, end)
        result["name"] = data.get("name", "").strip()
        result["birth_date"] = birth.strftime("%d/%m/%Y")
        result["end_date"] = end.strftime("%d/%m/%Y")
        return jsonify(result)
    except (KeyError, TypeError, ValueError) as error:
        return jsonify({"error": str(error) or "Please enter valid dates."}), 400


@app.post("/api/export")
def export_route():
    data = request.get_json()
    file_type = data.get("type", "txt").lower()
    rows = [
        ("Name", data.get("name") or "—"),
        ("Date of birth", data.get("birth_date", "")),
        ("Calculate age on", data.get("end_date", "")),
        ("Age", f'{data["age"]["years"]} years, {data["age"]["months"]} months, {data["age"]["days"]} days'),
        ("Born on", data.get("born_on", "")),
        ("Next birthday", data.get("next_birthday", "")),
        ("Birthday date", data.get("birthday_date", "")),
    ]
    rows.extend((f"Total {key}", str(value)) for key, value in data["totals"].items())

    output = BytesIO()
    filename = "age-calculation"

    try:
        if file_type == "txt":
            output.write("AGE CALCULATOR\n\n".encode())
            output.write("\n".join(f"{key}: {value}" for key, value in rows).encode())
            mimetype, extension = "text/plain", "txt"
        elif file_type == "xlsx":
            from openpyxl import Workbook
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Age Calculation"
            sheet.append(["Details", "Result"])
            for row in rows:
                sheet.append(list(row))
            workbook.save(output)
            mimetype, extension = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
        elif file_type == "docx":
            from docx import Document
            document = Document()
            document.add_heading("Age Calculator", 0)
            for key, value in rows:
                document.add_paragraph(f"{key}: {value}")
            document.save(output)
            mimetype, extension = "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
        elif file_type == "pdf":
            from reportlab.pdfgen import canvas
            pdf = canvas.Canvas(output)
            pdf.setFont("Helvetica-Bold", 18)
            pdf.drawString(50, 790, "Age Calculator")
            pdf.setFont("Helvetica", 11)
            y = 750
            for key, value in rows:
                pdf.drawString(50, y, f"{key}: {value}")
                y -= 24
            pdf.save()
            mimetype, extension = "application/pdf", "pdf"
        else:
            return jsonify({"error": "Unsupported file format."}), 400
    except ImportError as error:
        return jsonify({"error": f"Install the export package first: {error}"}), 500

    output.seek(0)
    return send_file(
        output,
        mimetype=mimetype,
        as_attachment=True,
        download_name=f"{filename}.{extension}",
    )


if __name__ == "__main__":
    app.run(debug=True)